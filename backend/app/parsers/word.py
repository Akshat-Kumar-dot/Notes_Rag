from __future__ import annotations

import io

import docx

from app.parsers.base import Page, ParseError, ParseResult


def parse(data: bytes) -> ParseResult:
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception:
        # python-docx only handles the modern zip-based format. Legacy .doc is
        # a different binary format entirely and needs LibreOffice.
        raise ParseError(
            "This file couldn't be read. If it's an older .doc, save it as .docx first."
        ) from None

    blocks = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                blocks.append(" | ".join(cells))

    text = "\n\n".join(blocks)
    if not text:
        raise ParseError("This document has no readable text.")

    return ParseResult(
        text=text,
        pages=[Page(number=1, text=text)],
        page_count=1,
        pages_with_text=1,
        parser_name="python-docx",
        parser_version="1",
    )
