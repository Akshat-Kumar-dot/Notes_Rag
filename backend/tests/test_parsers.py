import codecs
import io

import pytest

from app.parsers import plaintext, word
from app.parsers.base import ParseError


def test_utf8_text():
    r = plaintext.parse(b"# Title\n\nBody text.")
    assert r.coverage == 1.0 and r.chars > 0


def test_cp1252_fallback():
    """Windows-authored .txt files are frequently cp1252, not utf-8."""
    assert plaintext.parse("café — dash".encode("cp1252")).text == "café — dash"


def test_empty_file_rejected():
    with pytest.raises(ParseError):
        plaintext.parse(b"   \n  ")


def test_docx_paragraphs_and_tables():
    import docx

    d = docx.Document()
    d.add_paragraph("Queue length penalty under-weights throughput.")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "col a"
    t.rows[0].cells[1].text = "col b"
    buf = io.BytesIO()
    d.save(buf)

    r = word.parse(buf.getvalue())
    assert "throughput" in r.text
    assert "col a | col b" in r.text


def test_legacy_doc_gives_actionable_error():
    with pytest.raises(ParseError, match="docx"):
        word.parse(b"\xd0\xcf\x11\xe0legacy-ole-header")


NOTE = "# Exam notes\n\nThe TCP handshake is SYN, SYN-ACK, ACK — in that order."


@pytest.mark.parametrize("encoding", ["utf-8-sig", "utf-16", "utf-16-be", "utf-32"])
def test_bom_encodings_decode_cleanly(encoding):
    """Notepad and PowerShell on Windows write UTF-8 with a BOM and UTF-16."""
    data = NOTE.encode(encoding)
    if encoding == "utf-16-be":
        data = codecs.BOM_UTF16_BE + data
    assert plaintext.parse(data).text == NOTE


def test_bom_does_not_hide_the_first_heading():
    from app.rag.chunking import chunk_pages

    r = plaintext.parse(NOTE.encode("utf-16"))
    chunks = chunk_pages(r.pages, size=400, overlap=50)
    assert chunks[0].heading == "Exam notes"


def test_nul_bytes_rejected_with_a_clear_message():
    # UTF-16 without a BOM: nothing says what it is, and it decodes to NULs.
    with pytest.raises(ParseError, match="UTF-8"):
        plaintext.parse(NOTE.encode("utf-16-le"))
